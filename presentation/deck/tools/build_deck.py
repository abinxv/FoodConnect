#!/usr/bin/env python3
"""Assemble the final .pptx and/or .pdf. Lives in deck/tools/. Run only when the user asks.

    .venv/bin/python tools/build_deck.py                   pptx + pdf into ../export/
    .venv/bin/python tools/build_deck.py --name Ocea_SIH   output file name (default: deck.json title)
    .venv/bin/python tools/build_deck.py --pdf raster      pdf from the review PNGs (default, no template)
    .venv/bin/python tools/build_deck.py --pdf soffice     pdf by converting the pptx with LibreOffice (default with a template)
    .venv/bin/python tools/build_deck.py --pdf none | --no-pptx
    .venv/bin/python tools/build_deck.py --chrome-only     template decks: filled chrome only -> <chrome_dir>/<id>.png previews
    .venv/bin/python tools/build_deck.py --slides s1,s3    only these slides, in this order

No template (deck.json "template": null): one full-bleed picture per slide from ../slides/<id>.png.

With a template, deck.json "template" says how to use it:
  {
    "path": "templates/SIH.pptx",           relative to the project root
    "map": {"s1": {"page": 1, "picture": false}, "s2": 2, "s3": 3},
                                            deck slide -> template page (1-based); a page used twice is duplicated;
                                            "picture": false = fill fields only (e.g. a title slide)
    "fill": {"*": {"Oval 5": "Ocea"}, "s2": {"Title 1": "Ocea: 3D ocean viewer"}},
                                            shape name (or "ph:<idx>") -> text; "*" applies to every output slide
    "delete_shapes": {"*": ["TextBox 3"]},  shapes to remove (pointer / instruction boxes)
    "picture_z": "back",                    content picture behind the chrome ("front" puts it on top)
    "keep_unmapped": false                  drop template pages that no deck slide uses
  }
  Content pictures come from out/png/<id>.zone.png (node tools/export.mjs <id> --zone) and are placed
  on the export band (deck.json "band"). Run tools/inspect_template.py first to get shape names.

Needs python-pptx and Pillow (pip install python-pptx pillow). LibreOffice for --pdf soffice / --chrome-only.
"""
import argparse
import copy
import json
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
DECK = HERE.parent
ROOT = DECK.parent

try:
    from pptx import Presentation
    from pptx.util import Emu, Inches
    from pptx.oxml.ns import qn
except ImportError:
    sys.exit("python-pptx missing: .venv/bin/pip install python-pptx pillow")

RNS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def soffice_path():
    for c in [os.environ.get("SOFFICE"), shutil.which("soffice"), shutil.which("libreoffice"),
              "/Applications/LibreOffice.app/Contents/MacOS/soffice", "/opt/homebrew/bin/soffice",
              r"C:\Program Files\LibreOffice\program\soffice.exe"]:
        if c and Path(c).exists():
            return c
    return None


def to_pdf(pptx: Path, out_dir: Path) -> Path:
    so = soffice_path()
    if not so:
        sys.exit("LibreOffice not found (needed for pptx -> pdf). Install it, set SOFFICE, or use --pdf raster.")
    profile = Path(tempfile.gettempdir()) / "deck_lo_profile"
    # A separate profile avoids the silent failure when LibreOffice is already open.
    cmd = [so, f"-env:UserInstallation=file://{profile.as_posix()}", "--headless", "--convert-to", "pdf",
           "--outdir", str(out_dir), str(pptx)]
    r = subprocess.run(cmd, capture_output=True, text=True, timeout=600)
    pdf = out_dir / (pptx.stem + ".pdf")
    if not pdf.exists():
        sys.exit(f"LibreOffice conversion failed:\n{r.stdout}\n{r.stderr}")
    return pdf


# ---------- slide list helpers ----------
def slide_ids(prs):
    return list(prs.slides._sldIdLst)


def delete_slide(prs, slide):
    for sid in slide_ids(prs):
        if prs.slides.part.related_part(sid.rId) is slide.part:
            prs.part.drop_rel(sid.rId)
            prs.slides._sldIdLst.remove(sid)
            return


def reorder(prs, slides_in_order):
    lst = prs.slides._sldIdLst
    by_part = {prs.part.related_part(s.rId): s for s in slide_ids(prs)}
    for s in slides_in_order:
        el = by_part[s.part]
        lst.remove(el)
        lst.append(el)


def duplicate_slide(prs, src):
    """Copy a slide (shapes, background, pictures). Appended at the end."""
    new = prs.slides.add_slide(src.slide_layout)
    for shp in list(new.shapes):
        shp._element.getparent().remove(shp._element)
    rid_map = {}
    for rid, rel in src.part.rels.items():
        if "notesSlide" in rel.reltype or "slideLayout" in rel.reltype:
            continue
        if rel.is_external:
            rid_map[rid] = new.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        else:
            rid_map[rid] = new.part.relate_to(rel.target_part, rel.reltype)
    src_csld = src._element.find(qn("p:cSld"))
    new_csld = new._element.find(qn("p:cSld"))
    bg = src_csld.find(qn("p:bg"))
    if bg is not None:
        new_csld.insert(0, copy.deepcopy(bg))
    tree = new.shapes._spTree
    for el in src.shapes._spTree.iterchildren():
        if el.tag in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
            continue
        tree.append(copy.deepcopy(el))
    for el in new._element.iter():
        for k, v in list(el.attrib.items()):
            if k.startswith("{%s}" % RNS) and v in rid_map:
                el.set(k, rid_map[v])
    return new


# ---------- shapes ----------
def iter_shapes(shapes):
    for s in shapes:
        yield s
        if s.shape_type == 6:  # group
            yield from iter_shapes(s.shapes)


def find_shape(slide, key):
    if key.startswith("ph:"):
        idx = int(key[3:])
        for s in slide.placeholders:
            if s.placeholder_format.idx == idx:
                return s
        return None
    for s in iter_shapes(slide.shapes):
        if s.name == key:
            return s
    return None


def set_text(shape, text):
    """Replace text, keeping the first run's formatting. '\n' starts a new paragraph."""
    tf = shape.text_frame
    lines = str(text).split("\n")
    p0 = tf.paragraphs[0]
    proto = copy.deepcopy(p0._p)
    runs = p0.runs
    if runs:
        runs[0].text = lines[0]
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        p0.text = lines[0]
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    prev = p0._p
    for line in lines[1:]:
        np_ = copy.deepcopy(proto)
        rs = np_.findall(qn("a:r"))
        for r in rs[1:]:
            np_.remove(r)
        if rs:
            rs[0].find(qn("a:t")).text = line
        prev.addnext(np_)
        prev = np_


def add_picture(slide, img: Path, left, top, width, height, z="back"):
    pic = slide.shapes.add_picture(str(img), left, top, width, height)
    if z == "back":
        tree = slide.shapes._spTree
        tree.remove(pic._element)
        tree.insert(2, pic._element)  # after nvGrpSpPr + grpSpPr = behind everything
    return pic


def add_notes(slide, sid):
    f = DECK / "notes" / f"{sid}.md"
    if f.exists() and f.read_text().strip():
        slide.notes_slide.notes_text_frame.text = f.read_text().strip()


# ---------- builds ----------
def build_plain(deck, ids, out_pptx: Path):
    W, H = deck["canvas"]["width"], deck["canvas"]["height"]
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Emu(int(Inches(13.333) * H / W))
    blank = prs.slide_layouts[6]
    missing = [i for i in ids if not (ROOT / "slides" / f"{i}.png").exists()]
    if missing:
        sys.exit(f"missing review PNGs: {', '.join(missing)}. Run: node tools/export.mjs {' '.join(missing)}")
    for i in ids:
        s = prs.slides.add_slide(blank)
        add_picture(s, ROOT / "slides" / f"{i}.png", 0, 0, prs.slide_width, prs.slide_height, z="front")
        add_notes(s, i)
    prs.save(out_pptx)
    return out_pptx


def build_template(deck, ids, out_pptx: Path, chrome_only=False):
    t = deck["template"]
    tpath = (ROOT / t["path"]).resolve()
    if not tpath.exists():
        sys.exit(f"template not found: {tpath}")
    prs = Presentation(str(tpath))
    originals = list(prs.slides)
    W = deck["canvas"]["width"]
    emu_per_px = prs.slide_width / W
    band = deck["band"]
    used = {}
    ordered = []
    for sid in ids:
        spec = t["map"].get(sid)
        if spec is None:
            sys.exit(f"deck.json template.map has no entry for {sid}")
        if isinstance(spec, int):
            spec = {"page": spec}
        page = spec["page"]
        if page < 1 or page > len(originals):
            sys.exit(f"{sid}: template has {len(originals)} pages, map says {page}")
        src = originals[page - 1]
        slide = duplicate_slide(prs, src) if page in used else src
        used[page] = True
        ordered.append((sid, slide, spec))
    if not t.get("keep_unmapped"):
        keep = {id(s) for _, s, _ in ordered}
        for s in originals:
            if id(s) not in keep:
                delete_slide(prs, s)
    reorder(prs, [s for _, s, _ in ordered])

    for sid, slide, spec in ordered:
        fills = {**t.get("fill", {}).get("*", {}), **t.get("fill", {}).get(sid, {})}
        for key, text in fills.items():
            shp = find_shape(slide, key)
            if shp is None or not shp.has_text_frame:
                print(f"warning: {sid}: no text shape '{key}'")
                continue
            set_text(shp, text)
        dels = t.get("delete_shapes", {}).get("*", []) + t.get("delete_shapes", {}).get(sid, [])
        for key in dels:
            shp = find_shape(slide, key)
            if shp is None:
                print(f"warning: {sid}: no shape '{key}' to delete")
                continue
            shp._element.getparent().remove(shp._element)
        if chrome_only or spec.get("picture") is False:
            continue
        img = DECK / "out" / "png" / f"{sid}.zone.png"
        if not img.exists():
            sys.exit(f"missing {img.relative_to(ROOT)}. Run: node tools/export.mjs {sid} --zone")
        add_picture(slide, img, Emu(int(band["x"] * emu_per_px)), Emu(int(band["y"] * emu_per_px)),
                    Emu(int(band["w"] * emu_per_px)), Emu(int(band["h"] * emu_per_px)), t.get("picture_z", "back"))
        add_notes(slide, sid)
    prs.save(out_pptx)
    return out_pptx


def raster_pdf(ids, out_pdf: Path, canvas_w):
    from PIL import Image
    pages = []
    for i in ids:
        p = ROOT / "slides" / f"{i}.png"
        if not p.exists():
            sys.exit(f"missing {p}. Run: node tools/export.mjs {i}")
        pages.append(Image.open(p).convert("RGB"))
    dpi = pages[0].width / 13.333
    pages[0].save(out_pdf, save_all=True, append_images=pages[1:], resolution=dpi, quality=92)
    return out_pdf


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--name")
    ap.add_argument("--pdf", choices=["raster", "soffice", "none"])
    ap.add_argument("--no-pptx", action="store_true")
    ap.add_argument("--chrome-only", action="store_true")
    ap.add_argument("--slides")
    a = ap.parse_args()

    deck = json.loads((DECK / "deck.json").read_text())
    ids = a.slides.split(",") if a.slides else deck["slides"]
    if not ids:
        sys.exit("deck.json has no slides")
    has_t = bool(deck.get("template"))
    name = a.name or "".join(c if c.isalnum() or c in "-_" else "_" for c in deck.get("title", "deck"))

    if a.chrome_only:
        if not has_t:
            sys.exit("--chrome-only needs a template in deck.json")
        work = DECK / "out" / "chrome_build"
        work.mkdir(parents=True, exist_ok=True)
        pptx = build_template(deck, ids, work / "chrome.pptx", chrome_only=True)
        pdf = to_pdf(pptx, work)
        sys.path.insert(0, str(HERE))
        from pdf_to_png import render
        cdir = DECK / deck["chrome_dir"]
        cdir.mkdir(parents=True, exist_ok=True)
        pngs = render(pdf, work / "page", width=deck["canvas"]["width"])
        for sid, png in zip(ids, pngs):
            shutil.copyfile(png, cdir / f"{sid}.png")
        print(f"chrome previews: {len(pngs)} pages -> {cdir.relative_to(ROOT)}/<id>.png")
        return

    out = ROOT / "export"
    out.mkdir(exist_ok=True)
    pdf_mode = a.pdf or ("soffice" if has_t else "raster")
    pptx = None
    if not a.no_pptx or pdf_mode == "soffice":
        pptx = out / f"{name}.pptx"
        (build_template if has_t else build_plain)(deck, ids, pptx)
        print(f"pptx: {pptx.relative_to(ROOT)} ({len(ids)} slides)")
    if pdf_mode == "raster":
        if has_t:
            print("note: raster pdf skips the template chrome; use --pdf soffice for the real thing")
        pdf = raster_pdf(ids, out / f"{name}.pdf", deck["canvas"]["width"])
        print(f"pdf:  {pdf.relative_to(ROOT)} ({pdf.stat().st_size / 1e6:.1f} MB)")
    elif pdf_mode == "soffice":
        pdf = to_pdf(pptx, out)
        print(f"pdf:  {pdf.relative_to(ROOT)} ({pdf.stat().st_size / 1e6:.1f} MB)")
    if a.no_pptx and pptx and pdf_mode == "soffice":
        pptx.unlink()


if __name__ == "__main__":
    main()
